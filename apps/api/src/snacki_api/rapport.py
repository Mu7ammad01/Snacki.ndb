"""Rapport du pilotage en Excel (v1.1, réunion 5, demande 2).

Mêmes chiffres que l'écran de pilotage (pilotage.summary), pour la période choisie, avec la
date de génération. Le PDF, lui, est la page rapport du site enregistrée par le navigateur.
Les textes venant de la base (noms de produits) sont écrits comme du texte, jamais comme des
formules : un nom commençant par « = » ne peut pas s'exécuter dans Excel.
"""

from __future__ import annotations

from datetime import UTC, date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

PAY = {"cash": "Espèces", "bankily": "Bankily", "sedad": "Sedad", "bimbank": "Bimbank",
       "bamis": "Bamis"}  # fmt: skip
BOLD = Font(bold=True)


def _text(value: object) -> object:
    """Neutralise l'injection de formule (CSV/Excel) : texte commençant par = + - @."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list[object]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = BOLD
    for row in rows:
        ws.append([_text(v) for v in row])
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(width, 40)


def build(summary: dict, now: datetime | None = None) -> bytes:
    now = now or datetime.now(UTC)
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"
    start: date = summary["start"]
    end: date = summary["end"]
    lines: list[tuple[str, object]] = [
        ("Snacki · rapport de pilotage", ""),
        ("Période", f"du {start:%d/%m/%Y} au {end:%d/%m/%Y}"),
        ("Généré le", f"{now:%d/%m/%Y à %H:%M} (heure de Nouadhibou)"),
        ("", ""),
        ("Chiffre d'affaires (MRU)", summary["revenue_mru"]),
        ("Commandes", summary["orders"]),
        ("Panier moyen (MRU)", summary["average_basket_mru"]),
        ("Cumul depuis l'ouverture (MRU)", summary["all_time_mru"]),
        ("Tampons de fidélité", summary["loyalty"]["stamps"]),
        ("Commandes offertes", summary["loyalty"]["rewards"]),
        ("Montant offert (MRU)", summary["loyalty"]["discount_mru"]),
    ]
    for label, value in lines:
        ws.append([label, value])
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 36
    _sheet(
        wb,
        "Par jour",
        ["Jour", "App (MRU)", "Comptoir (MRU)", "Historique (MRU)", "Total (MRU)", "Commandes"],
        [
            [
                p["day"],
                p["app_mru"],
                p["comptoir_mru"],
                p["historique_mru"],
                p["app_mru"] + p["comptoir_mru"] + p["historique_mru"],
                p["orders"],
            ]
            for p in summary["by_day"]
        ],
    )
    _sheet(
        wb,
        "Produits",
        ["Produit", "Quantité", "Montant app et comptoir (MRU)"],
        [[t["label"], t["quantity"], t["revenue_mru"]] for t in summary["top"]],
    )
    _sheet(
        wb,
        "Paiements",
        ["Moyen", "Encaissements", "Montant (MRU)"],
        [
            [PAY.get(p["method"], p["method"]), p["count"], p["amount_mru"]]
            for p in summary["payments"]
        ],
    )
    for row in wb["Par jour"].iter_rows(min_row=2, max_col=1):
        row[0].number_format = "DD/MM/YYYY"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
