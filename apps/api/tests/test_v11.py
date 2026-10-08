"""v1.1 (réunion 5) : historique filtrable, rapport Excel, commandes par jour."""

from datetime import UTC, datetime
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import SecretStr
from sqlalchemy import select, text

from snacki_api import auth, rapport
from snacki_api.db import get_sessionmaker
from snacki_api.models import AuditLog, StaffRole, StaffUser

SECRET = "k" * 48
TODAY = datetime.now(UTC).date()


@pytest.fixture(autouse=True)
def base_vide():
    with get_sessionmaker()() as s:
        s.execute(
            text(
                "TRUNCATE history_item, history_sale, audit_log, staff_user, order_line, orders "
                "RESTART IDENTITY CASCADE"
            )
        )
        s.commit()


@pytest.fixture
def api(settings) -> TestClient:
    from snacki_api.main import create_app

    configured = settings.model_copy(
        update={
            "oauth_client_id": "client-test",
            "oauth_client_secret": SecretStr("secret-de-test"),
            "oauth_redirect_uri": "https://web.test/auth/callback",
            "session_secret": SecretStr(SECRET),
        }
    )
    return TestClient(create_app(configured))


def membre(role: StaffRole, name: str | None = None) -> dict[str, str]:
    with get_sessionmaker()() as s:
        user = StaffUser(email=f"{role.value}@gmail.com", role=role, display_name=name)
        s.add(user)
        s.commit()
        return {"X-Staff-Session": auth.new_session(SECRET, user.id, user.session_version, 8)}


def vendre(api, h, paid=True):
    payload = {"items": [{"product_id": "salade", "quantity": 2}]}
    if paid:
        payload["paid_method"] = "cash"
    r = api.post("/v1/caisse/orders", json=payload, headers=h)
    assert r.status_code == 201, r.text


def test_historique_reserve_et_filtrable(api):
    caissier = membre(StaffRole.CAISSIER, "Fatou")
    gerante = membre(StaffRole.GERANTE, "Mariem")
    vendre(api, caissier)
    assert api.get("/v1/historique", headers=caissier).status_code == 403

    d = api.get("/v1/historique", headers=gerante).json()
    labels = {(e["actor"], e["label"]) for e in d["entries"]}
    assert ("Fatou", "Vente au comptoir") in labels
    assert ("Fatou", "Accès refusé") in labels  # le refus de la caissière est lui-même tracé
    assert {p["name"] for p in d["people"]} == {"Fatou", "Mariem"}
    assert d["truncated"] is False

    fatou = next(p["id"] for p in d["people"] if p["name"] == "Fatou")
    only = api.get("/v1/historique", params={"person": fatou, "group": "refus"}, headers=gerante)
    assert [e["label"] for e in only.json()["entries"]] == ["Accès refusé"]


@pytest.mark.parametrize(
    "params",
    [
        {"group": "inconnu"},
        {"group": "DROP TABLE"},
        {"person": 0},
        {"start": "2026-10-08", "end": "2026-01-01"},
    ],
)
def test_historique_filtres_invalides(api, params):
    assert (
        api.get("/v1/historique", params=params, headers=membre(StaffRole.ADMIN)).status_code == 422
    )


def test_rapport_excel(api):
    h = membre(StaffRole.GERANTE)
    vendre(api, h)
    vendre(api, h)
    r = api.get("/v1/pilotage/rapport.xlsx", headers=h)
    assert r.status_code == 200 and r.headers["cache-control"] == "no-store"
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in r.headers["content-disposition"]
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Synthèse", "Par jour", "Produits", "Paiements"]
    synthese = {row[0]: row[1] for row in wb["Synthèse"].iter_rows(values_only=True)}
    assert synthese["Chiffre d'affaires (MRU)"] == 400 and synthese["Commandes"] == 2
    assert str(synthese["Généré le"]).startswith(datetime.now(UTC).strftime("%d/%m/%Y"))
    assert list(wb["Par jour"].iter_rows(values_only=True))[-1][-1] == 2  # commandes du jour
    with get_sessionmaker()() as s:
        assert s.scalar(select(AuditLog.action).where(AuditLog.action == "report")) == "report"


def test_rapport_reserve(api):
    assert api.get("/v1/pilotage/rapport.xlsx").status_code == 401
    r = api.get("/v1/pilotage/rapport.xlsx", headers=membre(StaffRole.CAISSIER))
    assert r.status_code == 403


@pytest.mark.parametrize("value", ['=HYPERLINK("x")', "+1", "-1", "@SUM(A1)"])
def test_rapport_sans_injection_de_formule(value):
    assert rapport._text(value) == "'" + value
    assert rapport._text("Salade") == "Salade" and rapport._text(12) == 12


def test_commandes_par_jour_dans_le_pilotage(api):
    h = membre(StaffRole.GERANTE)
    vendre(api, h)
    vendre(api, h, paid=False)  # non encaissée : exclue
    point = api.get("/v1/pilotage", headers=h).json()["by_day"][-1]
    assert point["day"] == TODAY.isoformat() and point["orders"] == 1
