"""J16 : achats, fournisseurs et stock (gérante et admin), dépenses et marge brute au pilotage."""

from datetime import UTC, datetime, timedelta
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import SecretStr
from sqlalchemy import select, text

from snacki_api import auth
from snacki_api.db import get_sessionmaker
from snacki_api.models import AuditLog, StaffRole, StaffUser, StockMove

SECRET = "k" * 48
TODAY = datetime.now(UTC).date()


@pytest.fixture(autouse=True)
def base_vide():
    with get_sessionmaker()() as s:
        s.execute(
            text(
                "TRUNCATE stock_move, purchase_line, purchase, ingredient, supplier, audit_log, "
                "staff_user, order_line, orders, history_item, history_sale "
                "RESTART IDENTITY CASCADE"
            )
        )
        s.commit()


@pytest.fixture
def api(settings) -> TestClient:
    from snacki_api.main import create_app

    configured = settings.model_copy(
        update={
            "oauth_client_id": "client-test",
            "oauth_client_secret": SecretStr("secret-de-test"),
            "oauth_redirect_uri": "https://web.test/auth/callback",
            "session_secret": SecretStr(SECRET),
        }
    )
    return TestClient(create_app(configured))


def membre(role: StaffRole) -> dict[str, str]:
    with get_sessionmaker()() as s:
        user = StaffUser(email=f"{role.value}@gmail.com", role=role, display_name=role.value)
        s.add(user)
        s.commit()
        return {"X-Staff-Session": auth.new_session(SECRET, user.id, user.session_version, 8)}


@pytest.fixture
def gerante() -> dict[str, str]:
    return membre(StaffRole.GERANTE)


def ingredient(api, h, name="Avocat", unit="kg", alert=None) -> dict:
    r = api.post("/v1/achats/ingredients", json={"name": name, "unit": unit, "alert_qty": alert},
                 headers=h)  # fmt: skip
    assert r.status_code == 201, r.text
    return r.json()


def fournisseur(api, h, name="Marché central") -> dict:
    r = api.post("/v1/achats/fournisseurs", json={"name": name, "phone": "+222 36 12 34 56"},
                 headers=h)  # fmt: skip
    assert r.status_code == 201, r.text
    return r.json()


def acheter(api, h, lines, supplier_id=None, day=None, status=201) -> dict:
    payload = {
        "purchase_day": (day or TODAY).isoformat(),
        "supplier_id": supplier_id,
        "payment_method": "cash",
        "lines": lines,
    }
    r = api.post("/v1/achats", json=payload, headers=h)
    assert r.status_code == status, r.text
    return r.json()


def stock(api, h, ingredient_id) -> float:
    rows = api.get("/v1/achats/ingredients", headers=h).json()
    return next(i["stock_qty"] for i in rows if i["id"] == ingredient_id)


# --- Droits --------------------------------------------------------------------------------


@pytest.mark.parametrize(
    "method,path",
    [
        ("get", "/v1/achats"),
        ("get", "/v1/achats/fournisseurs"),
        ("get", "/v1/achats/ingredients"),
        ("post", "/v1/achats"),
        ("post", "/v1/achats/1/annuler"),
        ("post", "/v1/achats/ingredients/1/inventaire"),
    ],
)
def test_reserve_a_la_gerante_et_a_l_admin(api, method, path):
    caissier = membre(StaffRole.CAISSIER)
    assert getattr(api, method)(path).status_code == 401
    kwargs = {"json": {}} if method == "post" else {}
    assert getattr(api, method)(path, headers=caissier, **kwargs).status_code in (403, 422)
    if method == "get":
        assert api.get(path, headers=caissier).status_code == 403
        assert api.get(path, headers=membre(StaffRole.ADMIN)).status_code == 200


def test_admin_et_gerante_ont_les_memes_droits(api, gerante):
    admin = membre(StaffRole.ADMIN)
    avocat = ingredient(api, gerante)
    p = acheter(api, admin, [{"ingredient_id": avocat["id"], "quantity": 3, "amount_mru": 900}])
    r = api.post(f"/v1/achats/{p['id']}/annuler", json={"reason": "erreur de saisie"},
                 headers=gerante)  # fmt: skip
    assert r.status_code == 200 and r.json()["cancel_reason"] == "erreur de saisie"


# --- Achats et stock -------------------------------------------------------------------------


def test_achat_total_calcule_et_stock_augmente(api, gerante):
    avocat = ingredient(api, gerante, alert=2)
    marche = fournisseur(api, gerante)
    assert marche["phone"] == "36123456"
    p = acheter(
        api,
        gerante,
        [
            {"ingredient_id": avocat["id"], "quantity": "2.5", "amount_mru": 750},
            {
                "label": "Gobelets 50 cl",
                "category": "emballages",
                "quantity": 100,
                "amount_mru": 400,
            },
        ],  # fmt: skip
        supplier_id=marche["id"],
    )
    assert p["total_mru"] == 1150 and p["supplier"] == "Marché central"
    assert p["created_by"] == "gerante"
    assert [line["label"] for line in p["lines"]] == ["Avocat", "Gobelets 50 cl"]
    assert stock(api, gerante, avocat["id"]) == 2.5
    rows = api.get("/v1/achats/ingredients", headers=gerante).json()
    assert rows[0]["low"] is False


def test_le_total_ne_vient_jamais_du_client(api, gerante):
    line = {"label": "Gaz", "category": "energie", "quantity": 1, "amount_mru": 500}
    r = api.post(
        "/v1/achats",
        json={
            "purchase_day": TODAY.isoformat(),
            "payment_method": "cash",
            "lines": [line],
            "total_mru": 1,
        },  # fmt: skip
        headers=gerante,
    )
    assert r.status_code == 422


@pytest.mark.parametrize(
    "line",
    [
        {"quantity": 1, "amount_mru": 10},  # ni ingrédient ni libellé
        {"ingredient_id": 1, "label": "x", "quantity": 1, "amount_mru": 10},  # les deux
        {"label": "Gaz", "category": "energie", "quantity": 0, "amount_mru": 10},
        {"label": "Gaz", "category": "energie", "quantity": 1, "amount_mru": 0},
        {"label": "Gaz", "category": "energie", "quantity": "1.0001", "amount_mru": 10},
        {"label": "Ga\u0000z", "category": "autre", "quantity": 1, "amount_mru": 10},
        {"ingredient_id": 1, "category": "energie", "quantity": 1, "amount_mru": 10},
    ],
)
def test_lignes_invalides_refusees(api, gerante, line):
    ingredient(api, gerante)
    acheter(api, gerante, [line], status=422)


def test_date_et_references_controlees(api, gerante):
    avocat = ingredient(api, gerante)
    line = {"ingredient_id": avocat["id"], "quantity": 1, "amount_mru": 10}
    acheter(api, gerante, [line], day=TODAY + timedelta(days=1), status=422)
    acheter(api, gerante, [line], day=TODAY - timedelta(days=400), status=422)
    acheter(api, gerante, [line, line], status=422)  # même ingrédient deux fois
    acheter(api, gerante, [{**line, "ingredient_id": 999}], status=409)
    acheter(api, gerante, [line], supplier_id=999, status=409)
    r = api.patch(f"/v1/achats/ingredients/{avocat['id']}", json={"active": False},
                  headers=gerante)  # fmt: skip
    assert r.status_code == 200
    acheter(api, gerante, [line], status=409)  # ingrédient désactivé


def test_annulation_reprend_le_stock_et_reste_visible(api, gerante):
    avocat = ingredient(api, gerante)
    p = acheter(api, gerante, [{"ingredient_id": avocat["id"], "quantity": 4, "amount_mru": 1200}])
    url = f"/v1/achats/{p['id']}/annuler"
    assert api.post(url, json={"reason": "x"}, headers=gerante).status_code == 422  # motif court
    r = api.post(url, json={"reason": "ticket saisi deux fois"}, headers=gerante)
    assert r.status_code == 200 and r.json()["cancelled_at"]
    assert stock(api, gerante, avocat["id"]) == 0
    again = api.post(url, json={"reason": "encore une fois"}, headers=gerante)
    assert again.status_code == 409
    d = api.get("/v1/achats", headers=gerante).json()
    assert len(d["purchases"]) == 1 and d["total_mru"] == 0  # visible, mais hors total


def test_inventaire_et_stock_jamais_negatif(api, gerante):
    avocat = ingredient(api, gerante, alert=1)
    p = acheter(api, gerante, [{"ingredient_id": avocat["id"], "quantity": 5, "amount_mru": 1500}])
    url = f"/v1/achats/ingredients/{avocat['id']}/inventaire"
    r = api.post(url, json={"quantity": "0.5", "reason": "comptage du soir"}, headers=gerante)
    assert r.status_code == 200 and r.json()["stock_qty"] == 0.5 and r.json()["low"] is True
    # Annuler l'achat retirerait 5 kg d'un stock de 0,5 kg : refusé, rien ne change.
    r = api.post(f"/v1/achats/{p['id']}/annuler", json={"reason": "erreur"}, headers=gerante)
    assert r.status_code == 409 and "négatif" in r.json()["detail"]
    assert stock(api, gerante, avocat["id"]) == 0.5
    moves = api.get(f"/v1/achats/ingredients/{avocat['id']}/mouvements", headers=gerante).json()
    assert [(m["kind"], m["delta"], m["stock_after"]) for m in moves] == [
        ("inventory", -4.5, 0.5),
        ("purchase", 5.0, 5.0),
    ]
    with get_sessionmaker()() as s:
        assert s.scalar(select(StockMove.reason).where(StockMove.kind == "inventory")) == (
            "comptage du soir"
        )


def test_noms_uniques_et_modifications(api, gerante):
    marche = fournisseur(api, gerante)
    r = api.post("/v1/achats/fournisseurs", json={"name": "Marché central"}, headers=gerante)
    assert r.status_code == 409
    r = api.patch(f"/v1/achats/fournisseurs/{marche['id']}", json={"phone": "", "active": False},
                  headers=gerante)  # fmt: skip
    assert r.status_code == 200 and r.json()["phone"] is None and r.json()["active"] is False
    bad = api.post("/v1/achats/fournisseurs", json={"name": "Boucherie", "phone": "12345678"},
                   headers=gerante)  # fmt: skip
    assert bad.status_code == 422
    ingredient(api, gerante)
    r = api.post("/v1/achats/ingredients", json={"name": "Avocat", "unit": "kg"}, headers=gerante)
    assert r.status_code == 409
    assert api.patch("/v1/achats/ingredients/999", json={}, headers=gerante).status_code == 404


def test_tout_est_journalise(api, gerante):
    avocat = ingredient(api, gerante)
    p = acheter(api, gerante, [{"ingredient_id": avocat["id"], "quantity": 1, "amount_mru": 300}])
    api.post(f"/v1/achats/{p['id']}/annuler", json={"reason": "erreur"}, headers=gerante)
    with get_sessionmaker()() as s:
        actions = s.scalars(select(AuditLog.action).order_by(AuditLog.id)).all()
    assert actions == ["ingredient_added", "purchase_added", "purchase_cancelled"]
    h = api.get("/v1/historique", params={"group": "achats"}, headers=gerante).json()
    assert {e["label"] for e in h["entries"]} == {
        "Ingrédient ajouté", "Achat enregistré", "Achat annulé"
    }  # fmt: skip


# --- Pilotage --------------------------------------------------------------------------------


def test_depenses_et_marge_brute(api, gerante):
    caissier = membre(StaffRole.CAISSIER)
    r = api.post(
        "/v1/caisse/orders",
        json={"items": [{"product_id": "salade", "quantity": 2}], "paid_method": "cash"},
        headers=caissier,
    )
    assert r.status_code == 201
    revenue = r.json()["total_mru"]
    avocat = ingredient(api, gerante, alert=10)
    marche = fournisseur(api, gerante)
    acheter(
        api,
        gerante,
        [{"ingredient_id": avocat["id"], "quantity": 2, "amount_mru": 600}],
        supplier_id=marche["id"],
    )
    acheter(api, gerante, [{"label": "Gaz", "category": "energie", "quantity": 1,
                            "amount_mru": 400}])  # fmt: skip
    d = api.get("/v1/pilotage", headers=gerante).json()
    assert d["purchases"]["total_mru"] == 1000
    assert d["gross_margin_mru"] == revenue - 1000
    assert d["purchases_first"] == TODAY.isoformat()
    cats = {c["category"]: c["amount_mru"] for c in d["purchases"]["by_category"]}
    assert cats == {"ingredients": 600, "energie": 400}
    sups = {x["supplier"]: x["amount_mru"] for x in d["purchases"]["by_supplier"]}
    assert sups == {"Marché central": 600, "Sans fournisseur": 400}
    assert [x["name"] for x in d["purchases"]["low_stock"]] == ["Avocat"]

    xlsx = api.get("/v1/pilotage/rapport.xlsx", headers=gerante)
    wb = load_workbook(BytesIO(xlsx.content))
    synthese = {row[0]: row[1] for row in wb["Synthèse"].iter_rows(values_only=True)}
    assert synthese["Achats (MRU)"] == 1000
    assert synthese["Marge brute (MRU) = ventes − achats"] == revenue - 1000
    achats_rows = list(wb["Achats"].iter_rows(values_only=True))
    assert ("Énergie (gaz, eau, électricité)", 400) in achats_rows


def test_nom_de_fournisseur_jamais_execute_dans_excel(api, gerante):
    piege = fournisseur(api, gerante, name='=HYPERLINK("http://x")')
    acheter(api, gerante, [{"label": "Sucre", "category": "autre", "quantity": 1,
                            "amount_mru": 50}], supplier_id=piege["id"])  # fmt: skip
    xlsx = api.get("/v1/pilotage/rapport.xlsx", headers=gerante)
    cells = [c for row in load_workbook(BytesIO(xlsx.content))["Achats"].iter_rows(values_only=True)
             for c in row]  # fmt: skip
    assert '\'=HYPERLINK("http://x")' in cells


def test_reponses_jamais_en_cache(api, gerante):
    assert api.get("/v1/achats", headers=gerante).headers["cache-control"] == "no-store"
