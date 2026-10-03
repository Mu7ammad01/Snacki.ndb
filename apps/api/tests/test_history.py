"""J8 : historique Excel — lecture, reconnaissance des articles, import idempotent (ADR 0010)."""

from datetime import date, datetime

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select, text

from snacki_api import history, manage
from snacki_api.db import get_sessionmaker
from snacki_api.models import HistoryItem, HistorySale


@pytest.fixture(autouse=True)
def base_vide():
    with get_sessionmaker()() as s:
        s.execute(text("TRUNCATE history_item, history_sale RESTART IDENTITY CASCADE"))
        s.commit()


def classeur(path, rows_by_day=None, sheet="Ventes", extra_sheet=True):
    """Reproduit la mise en page du « Suivi Dashboard » : un bloc de 5 colonnes par jour."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    rows_by_day = rows_by_day or {
        "J1 02/09/2026": [
            ("cmd1", "salade fruit", 2, 2000),
            ("cmd2", "2 crepes + 2 jus fraises", 4, 4400),
        ],
        "J2 03/09/2026": [("cmd1", "jus d'avocat", 1, 1000), ("cmd2", "teramisu", 1, 1500)],
    }
    ws.cell(4, 1, "S1")
    for k, (title, rows) in enumerate(rows_by_day.items()):
        col = 2 + 5 * k
        ws.cell(7, col, title)
        for j, h in enumerate(["Num cmd", "objet", "qtité", "prix"]):
            ws.cell(8, col + j, h)
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                ws.cell(9 + i, col + j, v)
    if extra_sheet:
        wb.create_sheet("New departs").cell(1, 1, "Capital investi — ne doit jamais être lu")
    wb.save(path)
    return path


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("salade fruit", [("salade", 1)]),
        ("2 crepes + 2 jus fraises", [("crepe", 2), ("fraise", 2)]),
        ("1 Asir mangue", [("mangue", 1)]),
        ("3 jus fraise + cokctail", [("fraise", 3), ("cocktail", 1)]),
        ("1 jus avocat (+ article illisible)", [("avocat", 1), (None, 1)]),
        ("2 gateaux ?", [(None, 2)]),
        ("Crêpe", [("crepe", 1)]),
        ("1 slade de fruits", [("salade", 1)]),
    ],
)
def test_reconnaissance_des_articles(raw, expected):
    assert [(i.product_id, i.quantity) for i in history.parse_items(raw)] == expected


def test_caracteres_de_controle_retires():
    (item,) = history.parse_items("salade\x00\nfruits")
    assert item.product_id == "salade" and "\x00" not in item.label


def test_lecture_du_classeur(tmp_path):
    sales, batch = history.read_workbook(classeur(tmp_path / "suivi.xlsx"))
    assert len(batch) == 16
    assert [(s.day, s.ref, s.total_mru) for s in sales] == [
        (date(2026, 9, 2), "cmd1", 200),  # 2000 MRO = 200 MRU
        (date(2026, 9, 2), "cmd2", 440),
        (date(2026, 9, 3), "cmd1", 100),
        (date(2026, 9, 3), "cmd2", 150),
    ]


def test_lignes_invalides_ignorees():
    grid = [
        ["J1 02/09/2026", None, None, None],
        ["Num cmd", "objet", "qtité", "prix"],
        ["cmd1", "salade", 1, 1000],
        ["cmd2", "salade", 1, None],  # pas de montant
        ["cmd3", None, 1, 1000],  # pas de texte
        ["total", "salade", 1, 1000],  # pas une commande
        ["cmd4", "salade", 1, -5],  # montant négatif
        ["cmd5", "salade", 1, True],  # booléen
    ]
    assert [s.ref for s in history.sales_from_grid(grid)] == ["cmd1"]


@pytest.mark.parametrize("name", ["suivi.csv", "suivi.xlsm"])
def test_extension_refusee(tmp_path, name):
    f = tmp_path / name
    f.write_text("x")
    with pytest.raises(history.HistoryError):
        history.read_workbook(f)


def test_fichier_trop_gros(tmp_path, monkeypatch):
    monkeypatch.setattr(history, "MAX_FILE_BYTES", 100)
    with pytest.raises(history.HistoryError, match="volumineux"):
        history.read_workbook(classeur(tmp_path / "suivi.xlsx"))


def test_feuille_ventes_obligatoire(tmp_path):
    with pytest.raises(history.HistoryError, match="Ventes"):
        history.read_workbook(classeur(tmp_path / "suivi.xlsx", sheet="Autre"))


def compte():
    with get_sessionmaker()() as s:
        return (
            s.scalar(select(func.count()).select_from(HistorySale)),
            s.scalar(select(func.sum(HistorySale.total_mru))),
            s.scalar(select(func.count()).select_from(HistoryItem)),
        )


def test_import_idempotent(tmp_path):
    sales, batch = history.read_workbook(classeur(tmp_path / "suivi.xlsx"))
    with get_sessionmaker()() as s:
        report = history.import_sales(s, sales, batch)
    assert (report.days, report.sales, report.total_mru) == (2, 4, 890)
    assert report.unknown == {"teramisu": 1}
    assert compte() == (4, 890, 5)
    with get_sessionmaker()() as s:
        history.import_sales(s, sales, batch)  # même fichier une 2e fois
    assert compte() == (4, 890, 5)


def test_import_remplace_seulement_les_jours_du_fichier(tmp_path):
    with get_sessionmaker()() as s:
        history.import_sales(s, history.read_workbook(classeur(tmp_path / "a.xlsx"))[0], "a")
    corrige = {"J2 03/09/2026": [("cmd1", "salade", 1, 1000)]}
    with get_sessionmaker()() as s:
        history.import_sales(
            s, history.read_workbook(classeur(tmp_path / "b.xlsx", corrige))[0], "b"
        )
    assert compte()[:2] == (3, 740)  # 02/09 gardé (640) + 03/09 remplacé (100)


def test_doublon_de_reference_ignore():
    sales = [
        history.Sale(date(2026, 9, 2), "cmd1", "salade", 100, history.parse_items("salade")),
        history.Sale(date(2026, 9, 2), "cmd1", "crepe", 120, history.parse_items("crepe")),
    ]
    with get_sessionmaker()() as s:
        assert history.import_sales(s, sales, "x").sales == 1


def test_commande_import_a_blanc_puis_reel(tmp_path, capsys):
    f = classeur(tmp_path / "suivi.xlsx")
    assert manage.main(["import-history", str(f), "--dry-run"]) == 0
    out = capsys.readouterr().out
    assert "4 ventes, 2 jours" in out and "890 MRU" in out and "rien n'a été écrit" in out
    assert compte()[0] == 0
    assert manage.main(["import-history", str(f)]) == 0
    assert compte()[0] == 4


def test_commande_import_refuse_un_mauvais_fichier(tmp_path, capsys):
    f = tmp_path / "suivi.csv"
    f.write_text("x")
    assert manage.main(["import-history", str(f)]) == 1
    assert "Refusé" in capsys.readouterr().err
    assert manage.main(["inconnue"]) == 2


def test_dates_au_format_excel():
    grid = [[datetime(2026, 9, 2)], ["J1 02/09/2026"], ["cmd1", "salade", 1, 1000]]
    assert len(history.sales_from_grid(grid)) == 1
